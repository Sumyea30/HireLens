import io
import json

from docx import Document
from docx.shared import Inches
from PIL import Image

from app.core.auth import hash_password
from app.core.config import Settings
from app.services.face import FaceService, HashFaceEngine
from app.services.intake import CVStorage, IntakeService
from app.services.photos import PhotoStorage

PDF = b"%PDF-1.4\n1 0 obj\n<<>>\nendobj\ntrailer\n<<>>\n%%EOF\n"


def _jpeg(color: tuple[int, int, int]) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (64, 64), color).save(buf, format="JPEG", quality=85)
    return buf.getvalue()


FACE_JPEG = _jpeg((180, 120, 90))
FACE_JPEG_B = _jpeg((40, 90, 200))


def make_docx(text="Backend developer with Python experience.", photo: bytes | None = None) -> bytes:
    buf = io.BytesIO()
    d = Document()
    d.add_paragraph(text)
    if photo:
        d.add_picture(io.BytesIO(photo), width=Inches(1.0))
    d.save(buf)
    return buf.getvalue()


def make_cv_docx(*, name="Jane Doe", email="jane@example.com", extra="", photo: bytes | None = FACE_JPEG) -> bytes:
    body = f"{name}\nEmail: {email}\n\nExperience\nPython developer with Docker experience. {extra}"
    return make_docx(body, photo=photo)


USERS = {
    "ivy": {"hash": hash_password("pw-ivy"), "role": "interviewer", "name": "Ivy"},
    "ada": {"hash": hash_password("pw-ada"), "role": "admin", "name": "Ada"},
    "guest": {"hash": hash_password("pw-guest"), "role": "candidate", "name": "Guest"},
}


def settings(**kw) -> Settings:
    defaults = {"face_provider": "hash", "photo_dir": "stored_photos"}
    defaults.update(kw)
    return Settings(auth_secret="test-secret", users_json=json.dumps(USERS), apply_email="hiring@example.com", **defaults)


def open_job(store, job_id="job-abc12345", owner_id="ivy"):
    store.create_job(job_id, owner_id)
    store.update_job(job_id, title="Backend Developer", profile={"title": "Backend Developer", "required_skills": ["Python", "Docker"],
                     "min_experience_years": 3, "qualifications": ["BSc"]})
    store.update_job(job_id, status="open")
    return job_id


def wired_intake(tmp_path, store, max_mb=10, on_queued=None, settings_obj=None):
    photos = PhotoStorage(str(tmp_path / "photos"))
    face = FaceService(HashFaceEngine(), threshold=0.35)
    s = settings_obj or settings(photo_dir=str(tmp_path / "photos"))
    return IntakeService(
        store, CVStorage(str(tmp_path / "cvs")), max_mb, on_queued=on_queued,
        photos=photos, face=face, settings=s,
    ), photos, face
